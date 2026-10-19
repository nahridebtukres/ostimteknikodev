import os
import platform
import socket
import subprocess
import threading
import time
from datetime import datetime

import paramiko
import psutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SSH_HOST, SSH_PORT, SSH_USER = "127.0.0.1", 2222, "labuser"
SSH_LAB_PASSWORD, SSH_WRONG_PASSWORD = "LabPass2222", "wrong_password_lab"
LAB_DIR, XLSX_FILE = "lab_auth", "ssh_monitor_log.xlsx"
INTERVAL, MAX_ROUNDS, SSH_TIMEOUT = 2, 10, 3
prev_bytes_recv = prev_bytes_sent = 0
prev_sample_time, failed_login_total = 0.0, 0

HEADERS = [
    "timestamp", "interface_name", "local_ip", "remote_ip", "remote_port", "local_port",
    "transport_protocol", "connection_state", "active_connection_count", "listening_port_count",
    "bytes_received", "bytes_sent", "packets_received", "packets_sent",
    "incoming_traffic_rate", "outgoing_traffic_rate", "ping_latency", "connection_setup_time",
    "cpu_utilization", "memory_usage", "failed_login_count",
]

def get_timestamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

def collect_interface_metrics():
    global prev_bytes_recv, prev_bytes_sent, prev_sample_time
    interface_name, local_ip = "unknown", "127.0.0.1"
    stats, addrs = psutil.net_if_stats(), psutil.net_if_addrs()
    ordered = [n for n, st in stats.items() if st.isup and not n.lower().startswith("lo")] + list(addrs)
    for name in ordered:
        for addr in addrs.get(name, []):
            if addr.family == socket.AF_INET:
                interface_name, local_ip = name, addr.address
                break
        if interface_name != "unknown":
            break
    io_all = psutil.net_io_counters()
    br, bs = int(io_all.bytes_recv), int(io_all.bytes_sent)
    pr, ps = int(io_all.packets_recv), int(io_all.packets_sent)
    now = time.time()
    dt = now - prev_sample_time if prev_sample_time > 0 else 0
    in_r = max(0.0, (br - prev_bytes_recv) / dt) if dt > 0 else 0.0
    out_r = max(0.0, (bs - prev_bytes_sent) / dt) if dt > 0 else 0.0
    prev_bytes_recv, prev_bytes_sent, prev_sample_time = br, bs, now
    return {
        "interface_name": interface_name, "local_ip": local_ip, "bytes_received": br, "bytes_sent": bs,
        "packets_received": pr, "packets_sent": ps,
        "incoming_traffic_rate": round(in_r, 2), "outgoing_traffic_rate": round(out_r, 2),
    }

def collect_connection_metrics():
    local_port, proto, state, active_n, listen_n = 0, "TCP", "NONE", 0, 0
    try:
        conns = psutil.net_connections(kind="inet")
    except (psutil.AccessDenied, PermissionError):
        conns = []
    for c in conns:
        if c.status == "LISTEN":
            listen_n += 1
        if c.status in ("ESTABLISHED", "SYN_SENT", "SYN_RECV", "TIME_WAIT"):
            active_n += 1
        if c.raddr and c.raddr.port == SSH_PORT and c.raddr.ip in (SSH_HOST, "127.0.0.1", "::1"):
            local_port = c.laddr.port if c.laddr else 0
            state = c.status
            proto = "UDP" if c.type == socket.SOCK_DGRAM else "TCP"
    return {
        "remote_ip": SSH_HOST, "remote_port": SSH_PORT, "local_port": local_port,
        "transport_protocol": proto, "connection_state": state,
        "active_connection_count": active_n, "listening_port_count": listen_n,
    }

def collect_system_metrics():
    cpu = float(psutil.cpu_percent(interval=0.2))
    mem = round(psutil.Process(os.getpid()).memory_info().rss / (1024 * 1024), 2)
    win = platform.system().lower() == "windows"
    cmd = ["ping", "-n" if win else "-c", "1", "-w" if win else "-W", "1000" if win else "1", SSH_HOST]
    ping = -1.0
    try:
        t0 = time.time()
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=3)
        if proc.returncode == 0:
            out = ((proc.stdout or "") + (proc.stderr or "")).lower()
            ping = round((time.time() - t0) * 1000.0, 2)
            if "time=" in out:
                tok = "".join(ch if ch.isdigit() or ch == "." else " " for ch in out.split("time=")[-1])
                ping = float(tok.split()[0])
    except Exception:
        ping = -1.0
    return {"cpu_utilization": cpu, "memory_usage": mem, "ping_latency": ping}

class LabSSH(paramiko.ServerInterface):
    def get_allowed_auths(self, username):
        return "password"

    def check_auth_password(self, username, password):
        try:
            for line in open(os.path.join(LAB_DIR, "users.txt"), encoding="utf-8"):
                line = line.strip()
                if line and not line.startswith("#") and ":" in line:
                    u, p = line.split(":", 1)
                    if username == u and password == p:
                        return paramiko.AUTH_SUCCESSFUL
        except OSError:
            pass
        return paramiko.AUTH_FAILED

    def check_channel_request(self, kind, chanid):
        return paramiko.OPEN_SUCCEEDED

def start_lab_ssh():
    os.makedirs(LAB_DIR, exist_ok=True)
    key_p, user_p = os.path.join(LAB_DIR, "host_key"), os.path.join(LAB_DIR, "users.txt")
    if not os.path.isfile(user_p):
        open(user_p, "w", encoding="utf-8").write("labuser:{0}\n".format(SSH_LAB_PASSWORD))
    key = paramiko.RSAKey(filename=key_p) if os.path.isfile(key_p) else paramiko.RSAKey.generate(2048)
    if not os.path.isfile(key_p):
        key.write_private_key_file(key_p)
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    sock.bind((SSH_HOST, SSH_PORT))
    sock.listen(8)

    def serve(client):
        tr = None
        try:
            tr = paramiko.Transport(client)
            tr.add_server_key(key)
            tr.start_server(server=LabSSH())
            end = time.time() + SSH_TIMEOUT
            while time.time() < end and tr.is_active():
                time.sleep(0.05)
        except Exception:
            pass
        finally:
            if tr is not None:
                tr.close()
            client.close()

    def loop():
        while True:
            threading.Thread(target=serve, args=(sock.accept()[0],), daemon=True).start()

    threading.Thread(target=loop, daemon=True).start()
    time.sleep(0.2)

def attempt_ssh_login():
    global failed_login_total
    setup, state = 0.0, "FAILED"
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    t0 = time.time()
    try:
        client.connect(
            hostname=SSH_HOST, port=SSH_PORT, username=SSH_USER,
            password=SSH_WRONG_PASSWORD, timeout=SSH_TIMEOUT,
            allow_agent=False, look_for_keys=False,
            banner_timeout=SSH_TIMEOUT, auth_timeout=SSH_TIMEOUT,
        )
        state = "ESTABLISHED"
    except paramiko.AuthenticationException:
        failed_login_total += 1; state = "AUTH_FAILED"
    except (paramiko.SSHException, socket.error, OSError, TimeoutError):
        failed_login_total += 1; state = "CONNECT_FAILED"
    finally:
        setup = (time.time() - t0) * 1000.0
        try:
            client.close()
        except Exception:
            pass
    return {"connection_setup_time": round(setup, 2), "failed_login_count": failed_login_total, "ssh_connection_state": state}

def export_xlsx(rows):
    if not rows:
        return
    hf, hn = PatternFill("solid", fgColor="17352B"), Font(bold=True, color="FFFFFF")
    yf = {"ARTIS": "DCFCE7", "AZALIS": "FFEDD5", "SABIT": "E5E7EB"}
    wb, first, last = Workbook(), rows[0], rows[-1]
    ws = wb.active
    ws.title = "ssh_monitor_log"
    for col, name in enumerate(HEADERS, 1):
        c = ws.cell(1, col, name); c.fill, c.font = hf, hn
        ws.column_dimensions[c.column_letter].width = 15
    for r_i, row in enumerate(rows, 2):
        for c_i, name in enumerate(HEADERS, 1):
            cell = ws.cell(r_i, c_i, row.get(name, ""))
            if name == "connection_state" and str(row.get(name, "")).upper() == "AUTH_FAILED":
                cell.fill = PatternFill("solid", fgColor="FFEDD5")
    ws2 = wb.create_sheet("Ozet")
    for col, name in enumerate(["Metrik", "Ilk", "Son", "Fark", "Yon", "Anlam"], 1):
        c = ws2.cell(1, col, name); c.fill, c.font = hf, hn
        ws2.column_dimensions[c.column_letter].width = 22
    keys = [("failed_login_count", "Basarisiz SSH giris birikimi"),
            ("active_connection_count", "Etkin TCP baglanti yogunlugu"),
            ("incoming_traffic_rate", "Gelen trafik hizi (B/s)"),
            ("outgoing_traffic_rate", "Giden trafik hizi (B/s)"),
            ("connection_setup_time", "SSH deneme suresi (ms)"),
            ("cpu_utilization", "CPU (%)"), ("memory_usage", "Bellek (MB)"), ("ping_latency", "Ping (ms)")]
    r_i = 2
    for key, meaning in keys:
        try:
            a, b = float(first[key]), float(last[key])
            yon = "ARTIS" if b > a else ("AZALIS" if b < a else "SABIT")
            for c_i, val in enumerate((key, a, b, round(abs(b - a), 2), yon, meaning), 1):
                cell = ws2.cell(r_i, c_i, val)
                if c_i == 5:
                    cell.fill, cell.font = PatternFill("solid", fgColor=yf[yon]), Font(bold=True)
            r_i += 1
        except (KeyError, TypeError, ValueError):
            pass
    auth_n = sum(1 for r in rows if str(r.get("connection_state", "")).upper() == "AUTH_FAILED")
    genel = "Tum turlarda AUTH_FAILED." if auth_n == len(rows) else ("Kismi AUTH_FAILED." if auth_n else "AUTH_FAILED yok.")
    ws2.cell(r_i + 1, 1, "GENEL_SONUC")
    ws2.cell(r_i + 1, 2, "Tur={0} AUTH_FAILED={1} Son_fail={2}".format(len(rows), auth_n, last.get("failed_login_count", 0)))
    ws2.cell(r_i + 2, 1, "Yorum"); ws2.cell(r_i + 2, 2, genel)
    wb.save(XLSX_FILE)

def main():
    start_lab_ssh()
    print("SSH Localhost Ag Izleme | {0}:{1} | tur={2}".format(SSH_HOST, SSH_PORT, MAX_ROUNDS))
    collected = []
    for n in range(1, MAX_ROUNDS + 1):
        try:
            row = {"timestamp": get_timestamp()}
            row.update(collect_interface_metrics())
            row.update(collect_connection_metrics())
            row.update(collect_system_metrics())
            ssh = attempt_ssh_login()
            row["connection_state"] = ssh["ssh_connection_state"]
            row["connection_setup_time"] = ssh["connection_setup_time"]
            row["failed_login_count"] = ssh["failed_login_count"]
            collected.append(row)
            print("[{0}] tur={1} state={2} fail={3}".format(row["timestamp"], n, row["connection_state"], row["failed_login_count"]))
        except Exception as exc:
            print("Tur {0} hata: {1}".format(n, exc))
        if n < MAX_ROUNDS:
            time.sleep(INTERVAL)
    try:
        export_xlsx(collected); print("Tamamlandi | {0}".format(XLSX_FILE))
    except PermissionError:
        print("XLSX kilitli, Excel'i kapatip tekrar deneyin.")

if __name__ == "__main__":
    main()
